"""Generování reportu do Word (.docx) — vedle PDF, stejný obsah/struktura.

Používá python-docx. Sdílí klasifikaci (build_findings) i i18n texty s HTML/PDF
reportem (``report_html``). Barvy závažnosti, hlavička s názvem laboratoře a
patička „SENSITIVE DATA" + čísla stran.
"""

import os

from docx import Document
from docx.shared import Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .report_classify import (
    SEVERITIES, SEVERITY_RANK, SEVERITY_COLOR, CVSS_BAND, OWASP_2025, section_title,
)
from .report_html import (
    T, results_subsections, OWASP_DESC, CVSS_NOTE, SEV_LABEL,
)

NAVY = RGBColor(0x16, 0x23, 0x3F)
RED = RGBColor(0xE2, 0x23, 0x1A)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _hexrgb(hexs):
    h = hexs.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _heading(doc, text, size=14, color=NAVY, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def _para(doc, text, size=10, color=None, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return p


def _chip_run(p, text, hexcolor):
    run = p.add_run(f" {text} ")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # podbarvení textu (highlight) přes w:shd na run
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    rpr.append(shd)


def _page_number_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText"); instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


_LAB_LINE = ("PT Lab · Penetration Testing Laboratory · Univerzita Tomáše Bati ve Zlíně, "
             "Fakulta aplikované informatiky",
             "PT Lab · Penetration Testing Laboratory · Tomas Bata University in Zlin, "
             "Faculty of Applied Informatics")
_SENSITIVE = ("! CITLIVÁ DATA – POUZE PRO AUTORIZOVANÉ POUŽITÍ !",
              "! SENSITIVE DATA – FOR AUTHORIZED USE ONLY !")


def _header_footer(doc, meta, lang):
    i = 1 if lang == "en" else 0
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(_LAB_LINE[i])
    r.bold = True; r.italic = True; r.font.size = Pt(8); r.font.color.rgb = NAVY
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run(_SENSITIVE[i] + "    ")
    rr.bold = True; rr.font.size = Pt(8); rr.font.color.rgb = RED
    _page_number_field(fp)


def _table(doc, headers, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        _set_cell_text(hdr[i], htext, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9)
        _shade_cell(hdr[i], "16233F")
    return t


def _risk_cell(cell, sev):
    _set_cell_text(cell, SEV_LABEL.get(sev, "?"), bold=True,
                   color=RGBColor(0xFF, 0xFF, 0xFF), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade_cell(cell, SEVERITY_COLOR.get(sev, "#666666"))


def _ip_of(target):
    t = target
    if "://" in t:
        t = t.split("://", 1)[1]
    return t.split(":")[0].split("/")[0]


def _port_of(target):
    t = target
    if "://" in t:
        t = t.split("://", 1)[1]
    return t.split(":")[1].split("/")[0] if ":" in t else ""


def build_docx(meta, result, options, out_path):
    """Vytvoří .docx report a uloží do ``out_path``. Vrací cestu."""
    options = options or {}
    lang = result.get("lang", options.get("lang", "cs"))
    rtype = options.get("report_type", "technical")
    findings = result.get("findings", [])
    summary = result.get("summary", {})
    human = options.get("human", {}) or {}

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    for s in doc.sections:
        s.top_margin = Mm(28); s.bottom_margin = Mm(18)
        s.left_margin = Mm(18); s.right_margin = Mm(18)
    _header_footer(doc, meta, lang)

    # --- Titulka ---
    logo = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                        "assets", "ptlab_logo.png")
    if os.path.exists(logo):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(logo, width=Mm(28))
        except Exception:
            pass
    sub = T("management_sub" if rtype == "management" else "technical_sub", lang)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub.upper()); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RED
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("title", "Pentest Report")); r.bold = True; r.font.size = Pt(28)
    r.font.color.rgb = NAVY
    for k, v in [(T("client", lang), meta.get("client", "—")),
                 (T("project", lang), meta.get("project_name", "—")),
                 (T("prepared_by", lang), meta.get("author", "—")),
                 (T("date", lang), meta.get("date", "—"))]:
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(f"{k}: "); rr.font.color.rgb = MUTED; rr.font.size = Pt(11)
        rv = pp.add_run(str(v)); rv.bold = True; rv.font.size = Pt(11)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = p.add_run("  " + T("confidential", lang) + "  ")
    rc.bold = True; rc.font.color.rgb = RED
    doc.add_page_break()

    # --- Executive summary ---
    if options.get("include_exec", True):
        _heading(doc, T("h_exec", lang), 16)
        verdict = ("Byly zjištěny kritické nálezy." if summary.get("CRITICAL") else
                   "Byly zjištěny nálezy vysoké závažnosti." if summary.get("HIGH") else
                   "Byly zjištěny nálezy střední/nízké závažnosti.")
        if lang == "en":
            verdict = ("Critical findings were identified." if summary.get("CRITICAL") else
                       "High-severity findings were identified." if summary.get("HIGH") else
                       "Medium/low-severity findings were identified.")
        _para(doc, f"{T('found_total', lang)}: {result.get('total', 0)}. {verdict}")
        t = _table(doc, [s for s in SEVERITIES])
        cells = t.add_row().cells
        for i, s in enumerate(SEVERITIES):
            _set_cell_text(cells[i], summary.get(s, 0), bold=True,
                           color=RGBColor(0xFF, 0xFF, 0xFF), size=14,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade_cell(cells[i], SEVERITY_COLOR[s])

    # --- Metodika ---
    if options.get("include_methodology", True):
        _heading(doc, T("h_owasp", lang), 14)
        used = set(result.get("owasp", {}).keys())
        t = _table(doc, [T("col_cat", lang), T("col_name", lang)])
        for code, name in OWASP_2025.items():
            c = t.add_row().cells
            _set_cell_text(c[0], code, bold=True, size=9)
            mark = " ✓" if code in used else ""
            _set_cell_text(c[1], f"{name}{mark} — {OWASP_DESC[code][1 if lang=='en' else 0]}", size=9)
        _heading(doc, T("h_scale", lang), 14)
        t = _table(doc, [T("col_risk", lang), T("col_name", lang), T("col_score", lang), T("col_desc", lang)])
        for s in SEVERITIES:
            c = t.add_row().cells
            _risk_cell(c[0], s)
            _set_cell_text(c[1], s, bold=True, size=9)
            _set_cell_text(c[2], CVSS_BAND[s], size=9)
            _set_cell_text(c[3], CVSS_NOTE[s][1 if lang == "en" else 0], size=9)

    # --- Cíle ---
    _heading(doc, T("h_objectives", lang), 14)
    _para(doc, human.get("scope") or T("scope_default", lang))
    scan_results = options.get("scan_results", {}) or {}
    ips = sorted((scan_results.get("tcp", {}) or {}).keys())
    if ips:
        _heading(doc, T("h_targets", lang), 12)
        for ip in ips:
            doc.add_paragraph(ip, style="List Bullet")

    # --- Nástroje ---
    tools = options.get("tools") or []
    if tools:
        _heading(doc, T("h_tools", lang), 12)
        t = _table(doc, [T("col_tool", lang), T("col_version", lang), T("col_purpose", lang)])
        for tl in tools:
            c = t.add_row().cells
            _set_cell_text(c[0], tl.get("name", ""), size=9)
            _set_cell_text(c[1], tl.get("version", "—"), size=9)
            _set_cell_text(c[2], tl.get("purpose", ""), size=9)

    # --- Výsledky dle typu testu -> cíl -> závažnost ---
    doc.add_page_break()
    _heading(doc, T("h_results", lang), 16)
    if not findings:
        _para(doc, T("no_findings", lang), color=MUTED)
    else:
        for sec, heading, secf in results_subsections(result, lang):
            _heading(doc, heading, 13, color=NAVY)
            by_ip = {}
            for f in secf:
                by_ip.setdefault(_ip_of(f["target"]), []).append(f)
            for ip in sorted(by_ip.keys()):
                _heading(doc, ip, 11, color=NAVY, space_before=6)
                t = _table(doc, [T("col_ports", lang), T("col_vuln", lang), T("col_risk", lang)])
                for f in sorted(by_ip[ip], key=lambda x: SEVERITY_RANK[x["severity"]]):
                    c = t.add_row().cells
                    _set_cell_text(c[0], _port_of(f["target"]) or "—", size=9)
                    _set_cell_text(c[1], f"{f.get('owasp','')} {f['title']}", size=9)
                    _risk_cell(c[2], f["severity"])
            # detail HIGH/CRITICAL
            for f in sorted(secf, key=lambda x: SEVERITY_RANK[x["severity"]]):
                if f["severity"] not in ("CRITICAL", "HIGH"):
                    continue
                p = doc.add_paragraph()
                rr = p.add_run(f"{f['id']} — {f['title']}")
                rr.bold = True; rr.font.color.rgb = NAVY; rr.font.size = Pt(11)
                _chip_run(p, f["severity"], SEVERITY_COLOR[f["severity"]])
                _para(doc, f"{f.get('owasp','')} {f.get('owasp_name','')} · {f['target']}",
                      size=9, color=MUTED)
                _para(doc, f["description"], size=10)
                if f.get("impact"):
                    _para(doc, f"{T('impact', lang)}: {f['impact']}", size=9)
                if options.get("include_recommendations", True) and f.get("recommendation"):
                    _para(doc, f"{T('recommendation', lang)}: {f['recommendation']}", size=9)
                if f.get("comment"):
                    _para(doc, f"{T('comment', lang)}: {f['comment']}", size=9, italic=True)

    # --- Shrnutí / závěr ---
    if rtype != "management":
        _heading(doc, T("h_summary", lang), 14)
        _para(doc, human.get("summary") or T("summary_default", lang))
        _heading(doc, T("h_conclusion", lang), 14)
        _para(doc, human.get("conclusion") or T("conclusion_default", lang))
    else:
        _heading(doc, T("h_recommendation", lang), 14)
        _para(doc, human.get("recommendation") or T("recommendation_default", lang))

    doc.save(out_path)
    return out_path
